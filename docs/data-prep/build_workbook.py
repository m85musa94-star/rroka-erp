import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as L

F = "Arial"
H_FILL = PatternFill("solid", fgColor="3A2A20")
IN_FILL = PatternFill("solid", fgColor="FFF7CC")
CALC_FILL = PatternFill("solid", fgColor="EDEDED")
EX_FONT = Font(name=F, italic=True, color="808080", size=10)
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
import os
ROWS = int(os.environ.get("ROWS", "500"))  # rows prepared for input in each table

wb = Workbook()

LISTS = {
    "yes_no": ["نعم", "لا"],
    "paid_from": ["البنك (حدد الحساب في العمود التالي)", "المالك من ماله الشخصي", "نقدًا من صندوق المنشأة", "عهدة موظف", "لم يُدفع بعد (آجل)"],
    "category": ["خامات ومواد مباشرة", "أجور عمالة مباشرة", "رواتب إدارية", "إيجار الورشة/المكتب", "كهرباء ومياه", "صيانة مكائن ومعدات",
                 "أدوات ومستهلكات ورشة", "وقود ونقل وشحن", "تركيب ومقاولو باطن", "رسوم حكومية وتأشيرات وإقامات", "تأمينات اجتماعية وطبي",
                 "اتصالات وإنترنت وبرامج", "تسويق وإعلان", "أتعاب مهنية (محاسب/محامٍ)", "رسوم بنكية", "مصاريف تأسيس", "ضيافة ونثريات مكتبية", "أخرى (وضّح في الملاحظات)"],
    "material_dest": ["مشروع مباشرة", "مخزون الورشة"],
    "owner_type": ["إيداع نقدي في حساب المنشأة", "دفع مصروف تأسيس من ماله", "دفع ثمن مكينة/أصل من ماله", "دفع مصروف تشغيلي من ماله", "سحب شخصي من حساب المنشأة"],
    "asset_state": ["جديد", "مستعمل"],
    "proj_status": ["عرض سعر", "قيد التنفيذ", "مُسلَّم", "متوقف", "ملغى"],
    "collect_method": ["تحويل بنكي", "نقدًا", "شيك", "نقاط بيع (مدى/بطاقة)"],
    "status": ["لم يبدأ", "جارٍ التجهيز", "مكتمل", "غير متوفر"],
    "entity": ["مؤسسة فردية", "شركة ذات مسؤولية محدودة", "شركة شخص واحد", "أخرى"],
    "vat_period": ["شهري", "ربع سنوي"],
}

def font(**k):
    return Font(name=F, size=k.pop("size", 10), **k)

def setup(ws, title, note):
    ws.sheet_view.rightToLeft = True
    ws["A1"] = title; ws["A1"].font = font(bold=True, size=14, color="3A2A20")
    ws["A2"] = note; ws["A2"].font = font(color="595959"); ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 48

def table(ws, cols, example, formulas=None, start=4):
    """cols: list of (header, width, list_name|None, number_format|None, is_formula). Row start=header."""
    formulas = formulas or {}
    ncol = len(cols)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=min(ncol, 10))
    for i, (h, w, lst, fmt, isf) in enumerate(cols, 1):
        c = ws.cell(start, i, h); c.font = font(bold=True, color="FFFFFF"); c.fill = H_FILL
        c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center"); c.border = BOX
        ws.column_dimensions[L(i)].width = w
    ws.row_dimensions[start].height = 36
    ws.freeze_panes = ws.cell(start + 1, 1)
    first, last = start + 1, start + ROWS
    for i, (h, w, lst, fmt, isf) in enumerate(cols, 1):
        col = L(i)
        if lst:
            dv = DataValidation(type="list", formula1=f"=قوائم!${list_col[lst]}$2:${list_col[lst]}${len(LISTS[lst])+1}", allow_blank=True,
                                showErrorMessage=True, errorTitle="قيمة غير مقبولة", error="اختر من القائمة المنسدلة")
            ws.add_data_validation(dv); dv.add(f"{col}{first}:{col}{last}")
        for r in range(first, last + 1):
            c = ws.cell(r, i)
            c.border = BOX; c.font = font()
            if fmt: c.number_format = fmt
            if isf:
                c.value = formulas[i].format(r=r); c.fill = CALC_FILL
            else:
                c.fill = IN_FILL
    # example row (first input row), clearly labelled
    for i, v in enumerate(example, 1):
        if v is not None and not cols[i-1][4]:
            ws.cell(first, i, v)
        ws.cell(first, i).font = EX_FONT
    ws.cell(start - 1, 1, "الصف الأول (بخط رمادي مائل) مثال توضيحي لطريقة الكتابة فقط — احذفه أو اكتب فوقه. الخلايا الصفراء للإدخال، والرمادية تُحسب تلقائيًا.").font = font(italic=True, color="9C5700")
    return first, last

# ---------- Lists sheet (created early so validations can reference it)
lists = wb.active; lists.title = "قوائم"; lists.sheet_view.rightToLeft = True
list_col = {}
for j, (k, vals) in enumerate(LISTS.items(), 1):
    list_col[k] = L(j)
    lists.cell(1, j, k).font = font(bold=True)
    for r, v in enumerate(vals, 2):
        lists.cell(r, j, v).font = font()
    lists.column_dimensions[L(j)].width = 34
lists.sheet_state = "hidden"

DATE = "yyyy-mm-dd"; MONEY = "#,##0.00;(#,##0.00);-"; PCT = "0.0%"

# ---------- 0. Guide
g = wb.create_sheet("دليل الاستخدام", 0); g.sheet_view.rightToLeft = True
g.column_dimensions["A"].width = 4; g.column_dimensions["B"].width = 120
lines = [
    ("إر روكا للأثاث — ملف تجهيز البيانات المالية من 2026-01-01", "title"),
    ("الغرض: جمع كل البيانات المالية منذ بداية النشاط بصورة صحيحة ومكتملة لإدخالها في المنظومة مرة واحدة، فتصبح المنظومة المرجع الوحيد للحسابات.", ""),
    ("", ""),
    ("قواعد التعبئة (إلزامية):", "h"),
    ("1. صف واحد لكل مستند (فاتورة، إيصال، حركة). لا تجمع عدة فواتير في صف واحد.", ""),
    ("2. التاريخ بالصيغة 2026-03-15 (سنة-شهر-يوم) وبالتقويم الميلادي.", ""),
    ("3. المبالغ أرقام فقط بلا كلمة «ريال» وبلا فواصل نصية. اكتب المبلغ قبل الضريبة والضريبة كلًّا في عموده كما هما في الفاتورة — لا تحسب الضريبة بنفسك.", ""),
    ("4. إن لم تكن فاتورة ضريبية (بلا رقم ضريبي للمورد أو بلا ضريبة): اكتب الضريبة 0 واختر «لا» في عمود الفاتورة الضريبية.", ""),
    ("5. لكل صف اسم ملف المرفق (صورة/PDF) كما هو محفوظ في المجلد، مثال: EXP-0001.pdf. سمِّ الملفات بالرقم التسلسلي في العمود الأول.", ""),
    ("6. لا تترك «دُفع من» فارغًا. إن لم تعرف المصدر اكتب ذلك في الملاحظات — الفارغ يتحول إلى استثناء للمراجعة.", ""),
    ("7. رمز المشروع يُكتب كما في ورقة «المشاريع» حرفيًا (مثال: P-001). المصروف العام يُترك رمز مشروعه فارغًا.", ""),
    ("8. لا تحذف صفًا مكتوبًا لتصحيحه؛ عدّله واكتب سبب التعديل في الملاحظات.", ""),
    ("9. لا تُدخل أي رقم تقديري على أنه فعلي. ما لا مستند له يُكتب في الملاحظات «بلا مستند» ويُراجع.", ""),
    ("", ""),
    ("الألوان:", "h"),
    ("أصفر = خلية إدخال.  رمادي = تُحسب تلقائيًا (لا تكتب فيها).  الصف الأول في كل ورقة رمادي مائل = مثال توضيحي يُحذف.", ""),
    ("", ""),
    ("ترتيب العمل المقترح:", "h"),
    ("أ) «قائمة المطلوبات» — ما يجب تجهيزه وحالته.   ب) «بيانات المنشأة» و«الحسابات البنكية».   ج) «المشاريع».   د) «المصروفات» و«مشتريات الخامات».", ""),
    ("هـ) «تمويل المالك» و«الأصول والمكائن».   و) «فواتير المبيعات» و«التحصيلات» (من دفترة).   ز) «الرواتب المدفوعة» و«إقرارات الضريبة».   ح) راجع «الملخص والفحوصات».", ""),
    ("", ""),
    ("ملفات ترفق مع هذا الملف في مجلد واحد:", "h"),
    ("كشوف البنوك بصيغة Excel أو CSV من بوابة البنك (لكل حساب من فتحه حتى اليوم) + صور/PDF كل الفواتير والإيصالات مسماة بأرقامها التسلسلية.", ""),
    ("", ""),
    ("ملاحظة: هذا الملف أداة تجهيز فقط وليس دفاتر محاسبية. القيود والأرصدة تُنشأ في المنظومة بعد الإدخال والمطابقة مع كشف البنك.", "warn"),
]
for i, (t, kind) in enumerate(lines, 1):
    c = g.cell(i, 2, t)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    c.font = font(bold=True, size=15, color="3A2A20") if kind == "title" else font(bold=True, size=11) if kind == "h" else font(color="9C0006", bold=True) if kind == "warn" else font(size=11)

# ---------- 1. Checklist
ck = wb.create_sheet("قائمة المطلوبات", 1)
setup(ck, "قائمة المطلوبات — ما يجب تجهيزه", "حدّث عمود «الحالة» أولًا بأول. الأولوية ١ = مطلوب لبدء الإدخال، ٢ = مطلوب لإكمال الحسابات، ٣ = مطلوب قبل إقفال السنة.")
items = [
    (1, "بيانات المنشأة", "السجل التجاري، نوع الكيان، رأس المال، الرقم الضريبي وتاريخ التسجيل، فترة الإقرار", "ورقة «بيانات المنشأة» + صورة السجل وشهادة التسجيل الضريبي"),
    (1, "كشوف كل الحسابات البنكية", "من تاريخ فتح الحساب حتى اليوم، لكل حساب", "Excel أو CSV من بوابة البنك (لا PDF فقط)"),
    (1, "قائمة الحسابات البنكية", "البنك، رقم الآيبان، رصيد الافتتاح ورصيد آخر الكشف", "ورقة «الحسابات البنكية»"),
    (1, "بيان المشاريع", "كل مشروع منذ 2026-01-01 مع العميل وقيمة العقد وحالته", "ورقة «المشاريع»"),
    (1, "فواتير المصروفات (مشاريع + عمومية)", "كل فاتورة وإيصال منذ 2026-01-01", "ورقة «المصروفات» + صورة/PDF لكل فاتورة"),
    (1, "فواتير مشتريات الخامات بالكميات", "الصنف والكمية والوحدة والسعر، ووجهة الخامة (مشروع/مخزون)", "ورقة «مشتريات الخامات»"),
    (2, "تمويل المالك ومصاريف التأسيس", "كل مبلغ دفعه المالك من ماله أو أودعه، وكل سحب شخصي", "ورقة «تمويل المالك» + الإيصالات"),
    (2, "المكائن والأصول", "كل مكينة ومعدة: تاريخ الشراء، البائع، السعر، جديد/مستعمل، ما يثبت الثمن", "ورقة «الأصول والمكائن» + الفواتير أو إقرار المالك بالقيمة"),
    (2, "فواتير المبيعات من دفترة", "كل فاتورة صدرت منذ 2026-01-01 (للاستيراد كسجل تاريخي)", "تصدير Excel من دفترة + ورقة «فواتير المبيعات»"),
    (2, "التحصيلات من العملاء", "كل مبلغ استُلم: التاريخ، العميل، الفاتورة، الحساب البنكي", "ورقة «التحصيلات»"),
    (2, "الرواتب المدفوعة", "لكل موظف وكل شهر: المبلغ المدفوع وتاريخه ومن أين دُفع", "ورقة «الرواتب المدفوعة» + مسيرات الرواتب إن وُجدت"),
    (2, "إقرارات ضريبة القيمة المضافة المقدمة", "كل إقرار قُدّم في 2026: الفترة والمبالغ وتاريخ السداد", "ورقة «إقرارات الضريبة» + نسخة الإقرار من بوابة الهيئة"),
    (2, "عقد إيجار الورشة", "قيمة الإيجار ومدته وطريقة الدفع", "صورة العقد"),
    (3, "قرار سياسة الإهلاك", "العمر الإنتاجي لكل نوع أصل (قرار المالك/المحاسب — لا يُفترض)", "يُبلَّغ كتابيًا"),
    (3, "اسم المحاسب المعتمد للقيود والإقرارات", "من يراجع ويعتمد القيود والإقرار الضريبي", "الاسم وبيانات التواصل"),
    (3, "التحقق من موجة الفوترة الإلكترونية (المرحلة الثانية)", "هل المنشأة مُبلَّغة بموجة ربط مع منصة «فاتورة» وموعدها", "إشعار الهيئة أو تأكيد من بوابة الهيئة"),
]
cols = [("الأولوية", 10, None, None, False), ("البند", 34, None, None, False), ("المطلوب بالتحديد", 60, None, None, False),
        ("الصيغة / أين يُجهَّز", 52, None, None, False), ("الحالة", 16, "status", None, False), ("المسؤول", 18, None, None, False), ("ملاحظات", 40, None, None, False)]
ck.merge_cells("A2:G2")
for i, (h, w, *_ ) in enumerate(cols, 1):
    c = ck.cell(4, i, h); c.font = font(bold=True, color="FFFFFF"); c.fill = H_FILL; c.border = BOX
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); ck.column_dimensions[L(i)].width = w
dv = DataValidation(type="list", formula1=f"=قوائم!${list_col['status']}$2:${list_col['status']}${len(LISTS['status'])+1}", allow_blank=True)
ck.add_data_validation(dv)
for r, (p, a, b, c_) in enumerate(items, 5):
    for i, v in enumerate([p, a, b, c_, "لم يبدأ", None, None], 1):
        c = ck.cell(r, i, v); c.border = BOX; c.font = font(bold=(i == 2)); c.alignment = Alignment(wrap_text=True, vertical="top")
        if i >= 5: c.fill = IN_FILL
    dv.add(f"E{r}")
ck.freeze_panes = "A5"
last_ck = 4 + len(items)

# ---------- 2. Company
co = wb.create_sheet("بيانات المنشأة", 2)
setup(co, "بيانات المنشأة", "تحدد هذه البيانات المعالجة المحاسبية والزكوية والضريبية. اكتب «غير متوفر» لما لا تعرفه — لا تتركه فارغًا.")
co.merge_cells("A2:C2")
fields = [("الاسم القانوني (كما في السجل)", None), ("نوع الكيان", "entity"), ("رقم السجل التجاري", None), ("تاريخ إصدار السجل", None),
          ("رأس المال المسجل (ريال)", None), ("المالك / الشركاء ونسبهم", None), ("الرقم الضريبي (15 رقمًا)", None), ("تاريخ التسجيل في ضريبة القيمة المضافة", None),
          ("فترة الإقرار الضريبي", "vat_period"), ("تاريخ بداية النشاط الفعلي", None), ("بداية السنة المالية", None), ("نهاية السنة المالية", None),
          ("هل للمنشأة نشاط أو حسابات قبل 2026-01-01؟", "yes_no"), ("المحاسب المعتمد (الاسم والجوال)", None), ("هل المنشأة مُبلَّغة بموجة ربط «فاتورة»؟ وموعدها", None)]
for i, (h, w) in enumerate([("البيان", 46), ("القيمة", 46), ("ملاحظات", 40)], 1):
    c = co.cell(4, i, h); c.font = font(bold=True, color="FFFFFF"); c.fill = H_FILL; c.border = BOX; co.column_dimensions[L(i)].width = w
for r, (f, lst) in enumerate(fields, 5):
    co.cell(r, 1, f).font = font(bold=True); co.cell(r, 1).border = BOX
    for cc in (2, 3):
        co.cell(r, cc).fill = IN_FILL; co.cell(r, cc).border = BOX; co.cell(r, cc).font = font()
    if lst:
        dv = DataValidation(type="list", formula1=f"=قوائم!${list_col[lst]}$2:${list_col[lst]}${len(LISTS[lst])+1}", allow_blank=True)
        co.add_data_validation(dv); dv.add(f"B{r}")
co.cell(8, 2).number_format = DATE; co.cell(12, 2).number_format = DATE; co.cell(14, 2).number_format = DATE

# ---------- 3. Banks
bk = wb.create_sheet("الحسابات البنكية", 3)
setup(bk, "الحسابات البنكية", "حساب في كل صف. الرصيد يُنقل من الكشف كما هو. اسم الحساب المختصر يُستخدم في الأوراق الأخرى في عمود «اسم الحساب البنكي».")
cols = [("اسم الحساب المختصر", 22, None, None, False), ("البنك", 20, None, None, False), ("رقم الآيبان", 30, None, "@", False), ("العملة", 10, None, None, False),
        ("تاريخ فتح الحساب", 14, None, DATE, False), ("الرصيد في 2026-01-01 (أو عند الفتح)", 18, None, MONEY, False), ("تاريخ آخر كشف", 14, None, DATE, False),
        ("الرصيد في آخر كشف", 18, None, MONEY, False), ("ملف الكشف (Excel/CSV) مرفق؟", 14, "yes_no", None, False), ("ملاحظات", 36, None, None, False)]
table(bk, cols, ["الراجحي-الجاري", "مصرف الراجحي", "SA00 0000 0000 0000 0000 0000", "SAR", "2026-01-05", 0, "2026-09-30", 12500.50, "نعم", None])

# ---------- 4. Projects
pj = wb.create_sheet("المشاريع", 4)
setup(pj, "المشاريع", "مشروع في كل صف منذ 2026-01-01. «رمز المشروع» قصير وفريد (P-001، P-002…) ويُستخدم في كل الأوراق الأخرى.")
cols = [("رمز المشروع", 12, None, None, False), ("اسم المشروع / الوصف", 32, None, None, False), ("العميل", 26, None, None, False), ("الرقم الضريبي للعميل", 18, None, "@", False),
        ("جوال العميل", 14, None, "@", False), ("قيمة العقد قبل الضريبة", 16, None, MONEY, False), ("ضريبة العقد", 14, None, MONEY, False),
        ("قيمة العقد شاملة الضريبة", 16, None, MONEY, True), ("تاريخ التعاقد/الاعتماد", 14, None, DATE, False), ("تاريخ التسليم", 14, None, DATE, False),
        ("الحالة", 14, "proj_status", None, False), ("رقم عرض السعر في دفترة", 16, None, "@", False), ("ملاحظات", 36, None, None, False)]
table(pj, cols, ["P-001", "مطبخ وخزائن فيلا — العليا", "مثال: مؤسسة النخبة", "300000000000003", "05XXXXXXXX", 40000, 6000, None, "2026-02-01", "2026-04-15", "مُسلَّم", "EST-0005", None],
      {8: "=IF(AND(F{r}=\"\",G{r}=\"\"),\"\",N(F{r})+N(G{r}))"})

# ---------- 5. Expenses
ex = wb.create_sheet("المصروفات", 5)
setup(ex, "المصروفات (مشاريع + عمومية)", "كل فاتورة أو إيصال مصروف منذ 2026-01-01 عدا الخامات بالكميات (لها ورقة مستقلة) والرواتب (لها ورقة مستقلة). المصروف على مشروع يُكتب رمزه؛ المصروف العام يُترك رمز مشروعه فارغًا.")
cols = [("الرقم التسلسلي", 12, None, None, False), ("تاريخ الفاتورة", 13, None, DATE, False), ("المورد", 24, None, None, False), ("الرقم الضريبي للمورد", 18, None, "@", False),
        ("رقم الفاتورة", 14, None, "@", False), ("فاتورة ضريبية؟", 11, "yes_no", None, False), ("الوصف", 30, None, None, False), ("التصنيف", 24, "category", None, False),
        ("رمز المشروع", 11, None, None, False), ("المبلغ قبل الضريبة", 14, None, MONEY, False), ("الضريبة", 12, None, MONEY, False), ("الإجمالي", 14, None, MONEY, True),
        ("نسبة الضريبة (فحص)", 11, None, PCT, True), ("دُفع من", 28, "paid_from", None, False), ("اسم الحساب البنكي / العهدة", 20, None, None, False), ("تاريخ الدفع", 13, None, DATE, False),
        ("مرجع الحركة في كشف البنك", 18, None, "@", False), ("اسم ملف المرفق", 16, None, None, False), ("ملاحظات", 32, None, None, False)]
table(ex, cols, ["EXP-0001", "2026-03-10", "مثال: مؤسسة الإضاءة الحديثة", "310000000000003", "INV-2231", "نعم", "إكسسوارات إضاءة للمطبخ", "خامات ومواد مباشرة",
                 "P-001", 1000, 150, None, None, "البنك (حدد الحساب في العمود التالي)", "الراجحي-الجاري", "2026-03-10", "FT26069XXXX", "EXP-0001.pdf", None],
      {12: "=IF(AND(J{r}=\"\",K{r}=\"\"),\"\",N(J{r})+N(K{r}))", 13: "=IF(N(J{r})=0,\"\",N(K{r})/J{r})"})

# ---------- 6. Material purchases
mp = wb.create_sheet("مشتريات الخامات", 6)
setup(mp, "مشتريات الخامات بالكميات", "صنف في كل صف (الفاتورة ذات ٣ أصناف = ٣ صفوف بنفس رقم الفاتورة). ضريبة الفاتورة تُكتب في صف الصنف الأول فقط. الوجهة: «مشروع مباشرة» إن صُرفت الخامة لمشروع بعينه، وإلا «مخزون الورشة».")
cols = [("الرقم التسلسلي للفاتورة", 14, None, None, False), ("تاريخ الفاتورة", 13, None, DATE, False), ("المورد", 24, None, None, False), ("الرقم الضريبي للمورد", 18, None, "@", False),
        ("رقم الفاتورة", 14, None, "@", False), ("الصنف / الخامة", 28, None, None, False), ("الوحدة", 10, None, None, False), ("الكمية", 10, None, "#,##0.###", False),
        ("سعر الوحدة قبل الضريبة", 14, None, MONEY, False), ("إجمالي الصنف قبل الضريبة", 15, None, MONEY, True), ("ضريبة الفاتورة (صف أول فقط)", 14, None, MONEY, False),
        ("الوجهة", 16, "material_dest", None, False), ("رمز المشروع", 11, None, None, False), ("دُفع من", 28, "paid_from", None, False),
        ("اسم الحساب البنكي / العهدة", 20, None, None, False), ("تاريخ الدفع", 13, None, DATE, False), ("اسم ملف المرفق", 16, None, None, False), ("ملاحظات", 30, None, None, False)]
table(mp, cols, ["PUR-0001", "2026-02-12", "مثال: مصنع الألواح", "300000000000003", "S-7781", "MDF 18 ملم", "لوح", 20, 95, None, 285, "مشروع مباشرة", "P-001",
                 "البنك (حدد الحساب في العمود التالي)", "الراجحي-الجاري", "2026-02-12", "PUR-0001.pdf", None],
      {10: "=IF(OR(H{r}=\"\",I{r}=\"\"),\"\",H{r}*I{r})"})

# ---------- 7. Owner
ow = wb.create_sheet("تمويل المالك", 7)
setup(ow, "تمويل المالك ومصاريف التأسيس والسحوبات", "كل مبلغ دفعه المالك من ماله للمنشأة أو أودعه في حسابها، وكل سحب شخصي من حساب المنشأة. المصروف الذي دفعه المالك يُكتب هنا وليس في ورقة المصروفات (لتجنّب التكرار).")
cols = [("الرقم التسلسلي", 12, None, None, False), ("التاريخ", 13, None, DATE, False), ("النوع", 30, "owner_type", None, False), ("الوصف", 34, None, None, False),
        ("المورد (إن وُجد)", 22, None, None, False), ("الرقم الضريبي للمورد", 18, None, "@", False), ("رقم الفاتورة/الإيصال", 14, None, "@", False),
        ("فاتورة ضريبية؟", 11, "yes_no", None, False), ("المبلغ قبل الضريبة", 14, None, MONEY, False), ("الضريبة", 12, None, MONEY, False), ("الإجمالي", 14, None, MONEY, True),
        ("التصنيف (للمصروف)", 24, "category", None, False), ("رمز المشروع", 11, None, None, False), ("اسم الحساب البنكي (للإيداع/السحب)", 20, None, None, False),
        ("اسم ملف المرفق", 16, None, None, False), ("ملاحظات", 32, None, None, False)]
table(ow, cols, ["OWN-0001", "2026-01-08", "دفع مصروف تأسيس من ماله", "رسوم إصدار السجل التجاري", "وزارة التجارة", None, "R-55102", "لا", 1200, 0, None,
                 "مصاريف تأسيس", None, None, "OWN-0001.pdf", None],
      {11: "=IF(AND(I{r}=\"\",J{r}=\"\"),\"\",N(I{r})+N(J{r}))"})

# ---------- 8. Assets
asx = wb.create_sheet("الأصول والمكائن", 8)
setup(asx, "الأصول والمكائن والمعدات", "كل مكينة أو معدة أو أصل تزيد قيمته على ما يُستهلك فورًا. المستعمل بلا فاتورة: اكتب القيمة المتفق عليها وما يثبتها (عقد بيع، تحويل بنكي، إقرار المالك). العمر الإنتاجي يُترك فارغًا حتى يُقرَّر.")
cols = [("رقم الأصل", 11, None, None, False), ("اسم الأصل", 28, None, None, False), ("الماركة / الموديل", 20, None, None, False), ("الرقم التسلسلي", 16, None, "@", False),
        ("جديد / مستعمل", 12, "asset_state", None, False), ("تاريخ الشراء", 13, None, DATE, False), ("تاريخ بدء التشغيل", 13, None, DATE, False), ("البائع", 22, None, None, False),
        ("فاتورة ضريبية؟", 11, "yes_no", None, False), ("الثمن قبل الضريبة", 14, None, MONEY, False), ("الضريبة", 12, None, MONEY, False), ("الإجمالي", 14, None, MONEY, True),
        ("دُفع من", 28, "paid_from", None, False), ("اسم الحساب البنكي", 18, None, None, False), ("ما يثبت الثمن", 26, None, None, False),
        ("العمر الإنتاجي المقترح (سنوات) — يُقرَّر لاحقًا", 16, None, "0", False), ("القدرة الكهربائية (كيلوواط) إن عُرفت", 14, None, "0.00", False),
        ("اسم ملف المرفق", 16, None, None, False), ("ملاحظات", 30, None, None, False)]
table(asx, cols, ["AST-001", "منشار طاولة", "مثال: SCM SI 400", "SN-XXXX", "مستعمل", "2026-01-20", "2026-01-25", "مثال: ورشة الأمانة", "لا", 18000, 0, None,
                  "المالك من ماله الشخصي", None, "عقد بيع + تحويل من حساب المالك", None, 5.5, "AST-001.pdf", None],
      {12: "=IF(AND(J{r}=\"\",K{r}=\"\"),\"\",N(J{r})+N(K{r}))"})

# ---------- 9. Sales invoices
si = wb.create_sheet("فواتير المبيعات", 9)
setup(si, "فواتير المبيعات (من دفترة)", "كل فاتورة مبيعات صدرت منذ 2026-01-01 كما هي في دفترة — تُستورد كسجل تاريخي ولا يُعاد إصدارها. يمكن لصق تصدير دفترة هنا وإضافة رمز المشروع.")
cols = [("رقم الفاتورة في دفترة", 16, None, "@", False), ("تاريخ الفاتورة", 13, None, DATE, False), ("العميل", 26, None, None, False), ("الرقم الضريبي للعميل", 18, None, "@", False),
        ("رمز المشروع", 11, None, None, False), ("نوع الفاتورة (دفعة مقدمة/مرحلية/نهائية)", 20, None, None, False), ("المبلغ قبل الضريبة", 14, None, MONEY, False),
        ("الضريبة", 12, None, MONEY, False), ("الإجمالي", 14, None, MONEY, True), ("إشعار دائن مرتبط (إن وُجد)", 16, None, "@", False), ("ملاحظات", 32, None, None, False)]
table(si, cols, ["INV-00012", "2026-02-03", "مثال: مؤسسة النخبة", "300000000000003", "P-001", "دفعة مقدمة", 20000, 3000, None, None, None],
      {9: "=IF(AND(G{r}=\"\",H{r}=\"\"),\"\",N(G{r})+N(H{r}))"})

# ---------- 10. Collections
cl = wb.create_sheet("التحصيلات", 10)
setup(cl, "التحصيلات من العملاء", "كل مبلغ استُلم من عميل. التحويل البنكي يُطابق لاحقًا مع كشف البنك تلقائيًا بالتاريخ والمبلغ والمرجع.")
cols = [("التاريخ", 13, None, DATE, False), ("العميل", 26, None, None, False), ("رمز المشروع", 11, None, None, False), ("رقم الفاتورة المسددة", 16, None, "@", False),
        ("المبلغ", 14, None, MONEY, False), ("طريقة التحصيل", 18, "collect_method", None, False), ("اسم الحساب البنكي المستلم", 20, None, None, False),
        ("مرجع الحركة في كشف البنك", 18, None, "@", False), ("رقم سند القبض (إن وُجد)", 14, None, "@", False), ("ملاحظات", 32, None, None, False)]
table(cl, cols, ["2026-02-04", "مثال: مؤسسة النخبة", "P-001", "INV-00012", 23000, "تحويل بنكي", "الراجحي-الجاري", "FT26035XXXX", None, None])

# ---------- 11. Payroll
pr = wb.create_sheet("الرواتب المدفوعة", 11)
setup(pr, "الرواتب المدفوعة", "صف لكل موظف لكل شهر بما دُفع فعلًا (المنظومة لا تحتسب الرواتب). العامل المباشر يُكتب رمز مشروعه إن عمل على مشروع واحد في الشهر، وإلا يُترك فارغًا.")
cols = [("الشهر (2026-01)", 12, None, "@", False), ("اسم الموظف", 24, None, None, False), ("المسمى", 18, None, None, False), ("عمالة مباشرة؟", 11, "yes_no", None, False),
        ("الراتب الأساسي", 13, None, MONEY, False), ("البدلات", 12, None, MONEY, False), ("خصومات/سلف", 12, None, MONEY, False), ("صافي المدفوع", 14, None, MONEY, True),
        ("تاريخ الدفع", 13, None, DATE, False), ("دُفع من", 28, "paid_from", None, False), ("اسم الحساب البنكي", 18, None, None, False), ("رمز المشروع", 11, None, None, False),
        ("ملاحظات", 30, None, None, False)]
table(pr, cols, ["2026-03", "مثال: خالد", "نجار", "نعم", 3000, 700, 0, None, "2026-03-28", "البنك (حدد الحساب في العمود التالي)", "الراجحي-الجاري", None, None],
      {8: "=IF(AND(E{r}=\"\",F{r}=\"\",G{r}=\"\"),\"\",N(E{r})+N(F{r})-N(G{r}))"})

# ---------- 12. VAT returns
vr = wb.create_sheet("إقرارات الضريبة", 12)
setup(vr, "إقرارات ضريبة القيمة المضافة المقدمة في 2026", "كما قُدّمت في بوابة الهيئة فعلًا. تُستخدم لمطابقة ضريبة المخرجات والمدخلات بعد إدخال البيانات واكتشاف أي فرق يحتاج تصحيحًا مع المستشار الضريبي.")
cols = [("الفترة من", 13, None, DATE, False), ("الفترة إلى", 13, None, DATE, False), ("المبيعات الخاضعة", 15, None, MONEY, False), ("ضريبة المخرجات", 14, None, MONEY, False),
        ("المشتريات الخاضعة", 15, None, MONEY, False), ("ضريبة المدخلات", 14, None, MONEY, False), ("صافي الضريبة", 14, None, MONEY, True), ("تاريخ التقديم", 13, None, DATE, False),
        ("تاريخ السداد", 13, None, DATE, False), ("المبلغ المسدد", 14, None, MONEY, False), ("ملاحظات", 32, None, None, False)]
table(vr, cols, ["2026-01-01", "2026-03-31", 100000, 15000, 0, 0, None, "2026-04-28", "2026-04-28", 15000, "مثال: لم تُدرج مدخلات الربع"],
      {7: "=IF(AND(D{r}=\"\",F{r}=\"\"),\"\",N(D{r})-N(F{r}))"})

# ---------- 13. Summary & checks
sm = wb.create_sheet("الملخص والفحوصات", 13); sm.sheet_view.rightToLeft = True
sm["A1"] = "الملخص والفحوصات (تُحسب تلقائيًا — بعد حذف الصفوف التوضيحية)"; sm["A1"].font = font(bold=True, size=14, color="3A2A20")
sm.column_dimensions["A"].width = 58; sm.column_dimensions["B"].width = 18; sm.column_dimensions["C"].width = 18; sm.column_dimensions["D"].width = 60
for i, h in enumerate(["البند", "العدد", "المبلغ", "التفسير"], 1):
    c = sm.cell(3, i, h); c.font = font(bold=True, color="FFFFFF"); c.fill = H_FILL; c.border = BOX
R = f"5:{4+ROWS}"
def rng(sheet, col): return f"'{sheet}'!{col}5:{col}{4+ROWS}"
rows = [
    ("المشاريع", f"=SUMPRODUCT(1*({rng('المشاريع','A')}<>\"\"))", f"=SUM({rng('المشاريع','F')})", "قيمة العقود قبل الضريبة"),
    ("المصروفات", f"=SUMPRODUCT(1*({rng('المصروفات','A')}<>\"\"))", f"=SUM({rng('المصروفات','J')})", "قبل الضريبة"),
    ("  منها على مشاريع", f"=SUMPRODUCT(({rng('المصروفات','A')}<>\"\")*({rng('المصروفات','I')}<>\"\"))", f"=SUMPRODUCT(({rng('المصروفات','I')}<>\"\")*{rng('المصروفات','J')})", ""),
    ("  منها عامة", f"=SUMPRODUCT(({rng('المصروفات','A')}<>\"\")*({rng('المصروفات','I')}=\"\"))", f"=SUMPRODUCT(({rng('المصروفات','A')}<>\"\")*({rng('المصروفات','I')}=\"\")*{rng('المصروفات','J')})", ""),
    ("ضريبة المدخلات في المصروفات", None, f"=SUM({rng('المصروفات','K')})", ""),
    ("مشتريات الخامات (أصناف)", f"=SUMPRODUCT(1*({rng('مشتريات الخامات','F')}<>\"\"))", f"=SUM({rng('مشتريات الخامات','J')})", "قبل الضريبة"),
    ("تمويل المالك (كل الأنواع عدا السحب)", f"=SUMPRODUCT(({rng('تمويل المالك','A')}<>\"\")*({rng('تمويل المالك','C')}<>\"سحب شخصي من حساب المنشأة\"))", f"=SUMPRODUCT(({rng('تمويل المالك','C')}<>\"سحب شخصي من حساب المنشأة\")*{rng('تمويل المالك','I')})", "قبل الضريبة"),
    ("سحوبات المالك", f"=SUMPRODUCT(1*({rng('تمويل المالك','C')}=\"سحب شخصي من حساب المنشأة\"))", f"=SUMPRODUCT(({rng('تمويل المالك','C')}=\"سحب شخصي من حساب المنشأة\")*{rng('تمويل المالك','I')})", ""),
    ("الأصول والمكائن", f"=SUMPRODUCT(1*({rng('الأصول والمكائن','A')}<>\"\"))", f"=SUM({rng('الأصول والمكائن','J')})", "الثمن قبل الضريبة"),
    ("فواتير المبيعات", f"=SUMPRODUCT(1*({rng('فواتير المبيعات','A')}<>\"\"))", f"=SUM({rng('فواتير المبيعات','G')})", "قبل الضريبة"),
    ("ضريبة المخرجات في فواتير المبيعات", None, f"=SUM({rng('فواتير المبيعات','H')})", ""),
    ("التحصيلات", f"=SUMPRODUCT(1*({rng('التحصيلات','A')}<>\"\"))", f"=SUM({rng('التحصيلات','E')})", ""),
    ("الرواتب المدفوعة", f"=SUMPRODUCT(1*({rng('الرواتب المدفوعة','B')}<>\"\"))", f"=SUM({rng('الرواتب المدفوعة','H')})", "صافي المدفوع"),
    ("", None, None, ""),
    ("فحوصات (يجب أن تكون صفرًا)", None, None, ""),
    ("مصروفات بلا «دُفع من»", f"=SUMPRODUCT(({rng('المصروفات','A')}<>\"\")*({rng('المصروفات','N')}=\"\"))", None, "كل مصروف يجب أن يُعرف مصدر دفعه"),
    ("مصروفات بلا اسم ملف مرفق", f"=SUMPRODUCT(({rng('المصروفات','A')}<>\"\")*({rng('المصروفات','R')}=\"\"))", None, "مصروف بلا مستند = استثناء"),
    ("مصروفات بلا تصنيف", f"=SUMPRODUCT(({rng('المصروفات','A')}<>\"\")*({rng('المصروفات','H')}=\"\"))", None, ""),
    ("مصروفات بضريبة دون فاتورة ضريبية", f"=SUMPRODUCT(({rng('المصروفات','K')}>0)*({rng('المصروفات','F')}=\"لا\"))", None, "الضريبة لا تُسترد بلا فاتورة ضريبية"),
    ("مصروفات برقم فاتورة مكرر لنفس المورد", f"=SUMPRODUCT(({rng('المصروفات','E')}<>\"\")*(COUNTIFS({rng('المصروفات','C')},{rng('المصروفات','C')},{rng('المصروفات','E')},{rng('المصروفات','E')})>1))", None, "احتمال إدخال مكرر"),
    ("مصروفات على رمز مشروع غير موجود في ورقة المشاريع", f"=SUMPRODUCT(({rng('المصروفات','I')}<>\"\")*(COUNTIF({rng('المشاريع','A')},{rng('المصروفات','I')})=0))", None, "اكتب الرمز كما في ورقة المشاريع"),
    ("فواتير مبيعات بلا رمز مشروع", f"=SUMPRODUCT(({rng('فواتير المبيعات','A')}<>\"\")*({rng('فواتير المبيعات','E')}=\"\"))", None, ""),
    ("كشوف بنوك غير مرفقة", f"=SUMPRODUCT(({rng('الحسابات البنكية','A')}<>\"\")*({rng('الحسابات البنكية','I')}<>\"نعم\"))", None, ""),
    ("بنود «قائمة المطلوبات» غير المكتملة", f"=SUMPRODUCT(1*('قائمة المطلوبات'!E5:E{last_ck}<>\"مكتمل\"))", None, ""),
]
for r, (a, b, c_, d) in enumerate(rows, 4):
    sm.cell(r, 1, a).font = font(bold=a.startswith("فحوصات"))
    for i, v in ((2, b), (3, c_)):
        cc = sm.cell(r, i, v); cc.font = font(); cc.number_format = MONEY if i == 3 else "#,##0"
        if v: cc.fill = CALC_FILL
    sm.cell(r, 4, d).font = font(color="595959")
    for i in range(1, 5): sm.cell(r, i).border = BOX if a else Border()

wb.move_sheet("قوائم", offset=len(wb.sheetnames))
wb.active = 0
wb.calculation.fullCalcOnLoad = True
out = sys.argv[1]; wb.save(out); print("saved", out)
